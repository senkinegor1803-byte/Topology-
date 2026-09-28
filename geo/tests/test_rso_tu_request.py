"""Тесты Шага 3.11: разбор полей заявки, заполнение DOCX, выкопировка PDF,
пакет для ручной подачи."""

from __future__ import annotations

import io
import json
import zipfile

import docx
from pypdf import PdfReader

from topology_geo.rso.tu_request import (
    TuRequestFields,
    build_default_template,
    build_site_excerpt_pdf,
    build_tu_request_package,
    fill_docx_template,
)


def _fields() -> TuRequestFields:
    return TuRequestFields(
        applicant="ООО «Застройщик»", cadastral_number="59:01:1234567:89",
        address="г. Пермь, ул. Тестовая, 1", object_purpose="многоквартирный жилой дом",
        loads={"В": "5 м³/сут", "Т": "0.5 Гкал/ч"}, deadlines="IV квартал 2026",
    )


def test_build_default_template_is_valid_docx_with_placeholders():
    template_bytes = build_default_template()

    doc = docx.Document(io.BytesIO(template_bytes))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "{{заявитель}}" in full_text
    assert "{{кадастровый_номер}}" in full_text


def test_fill_docx_template_replaces_all_placeholders():
    template_bytes = build_default_template()
    filled_bytes = fill_docx_template(template_bytes, _fields())

    doc = docx.Document(io.BytesIO(filled_bytes))
    full_text = "\n".join(p.text for p in doc.paragraphs)

    assert "ООО «Застройщик»" in full_text
    assert "59:01:1234567:89" in full_text
    assert "г. Пермь, ул. Тестовая, 1" in full_text
    assert "многоквартирный жилой дом" in full_text
    assert "В: 5 м³/сут" in full_text
    assert "IV квартал 2026" in full_text
    assert "{{" not in full_text  # ни один плейсхолдер не остался незаполненным


def test_fill_docx_template_handles_split_runs():
    """Реальная особенность python-docx: текст параграфа, добавленный по
    частям (`add_run` несколько раз), хранится в НЕСКОЛЬКИХ `run` - плейсхолдер
    может физически разорваться между ними. Тест это воспроизводит явно,
    не полагаясь на то, что `Document.add_paragraph(text)` всегда даёт один run."""
    doc = docx.Document()
    p = doc.add_paragraph()
    p.add_run("Заявитель: {{за")
    p.add_run("явитель}}")
    buf = io.BytesIO()
    doc.save(buf)

    filled_bytes = fill_docx_template(buf.getvalue(), _fields())

    filled_doc = docx.Document(io.BytesIO(filled_bytes))
    assert "ООО «Застройщик»" in filled_doc.paragraphs[0].text


def test_build_site_excerpt_pdf_contains_network_summary():
    pdf_bytes = build_site_excerpt_pdf(
        site_name="Участок №1", footprint_area_m2=987.6,
        network_summary=["К1: канализация, 2.1 м", "В1: водопровод, 5.0 м"],
    )

    assert pdf_bytes[:4] == b"%PDF"
    text = "".join(page.extract_text() for page in PdfReader(io.BytesIO(pdf_bytes)).pages)
    assert "Участок №1" in text
    assert "987.6" in text
    assert "канализация" in text


def test_build_tu_request_package_contains_all_files():
    fields = _fields()
    docx_bytes = fill_docx_template(build_default_template(), fields)
    pdf_bytes = build_site_excerpt_pdf(site_name="Участок", footprint_area_m2=100.0, network_summary=[])

    package_bytes = build_tu_request_package(fields=fields, docx_bytes=docx_bytes, excerpt_pdf_bytes=pdf_bytes)

    zf = zipfile.ZipFile(io.BytesIO(package_bytes))
    assert zf.testzip() is None
    assert set(zf.namelist()) == {"request.docx", "attachment.pdf", "fields.json"}

    fields_json = json.loads(zf.read("fields.json"))
    assert fields_json["заявитель"] == "ООО «Застройщик»"
    assert fields_json["нагрузки"]["В"] == "5 м³/сут"

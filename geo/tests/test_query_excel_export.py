"""Тесты Шага 3.5, п. 3: фильтр/поиск зон + выгрузка в Excel."""

from __future__ import annotations

import io
from datetime import datetime, timezone

import openpyxl
import pytest

from topology_geo.constraints.store import (
    ConstraintZone,
    STATUS_CALCULATED,
    STATUS_OFFICIAL,
    ensure_schema,
    load_zone,
    search_zones,
)
from topology_geo.query.excel_export import COLUMNS, zones_to_excel

TS = datetime(2026, 1, 1, tzinfo=timezone.utc)
SQUARE = {"type": "Polygon", "coordinates": [[
    [56.20, 58.00], [56.21, 58.00], [56.21, 58.01], [56.20, 58.01], [56.20, 58.00],
]]}


@pytest.fixture()
def db(pg_test_db):
    ensure_schema(pg_test_db)
    return pg_test_db


def test_search_zones_filters_by_type(db):
    load_zone(db, ConstraintZone(
        zone_type="ООПТ", status=STATUS_OFFICIAL, source_name="s", data_timestamp=TS,
        geom_geojson=SQUARE, registry_number="A1",
    ))
    load_zone(db, ConstraintZone(
        zone_type="охранная зона ЛЭП", status=STATUS_CALCULATED, source_name="s", data_timestamp=TS,
        geom_geojson=SQUARE,
    ))

    results = search_zones(db, zone_type="ООПТ")
    assert len(results) == 1
    assert results[0].zone_type == "ООПТ"


def test_search_zones_filters_by_status(db):
    load_zone(db, ConstraintZone(
        zone_type="ООПТ", status=STATUS_OFFICIAL, source_name="s", data_timestamp=TS,
        geom_geojson=SQUARE, registry_number="A1",
    ))
    load_zone(db, ConstraintZone(
        zone_type="охранная зона ЛЭП", status=STATUS_CALCULATED, source_name="s", data_timestamp=TS,
        geom_geojson=SQUARE,
    ))

    results = search_zones(db, status=STATUS_CALCULATED)
    assert len(results) == 1
    assert results[0].status == STATUS_CALCULATED


def test_search_zones_filters_by_registry_number_substring(db):
    load_zone(db, ConstraintZone(
        zone_type="ООПТ", status=STATUS_OFFICIAL, source_name="s", data_timestamp=TS,
        geom_geojson=SQUARE, registry_number="59:00-1.1",
    ))
    load_zone(db, ConstraintZone(
        zone_type="ОКН", status=STATUS_OFFICIAL, source_name="s", data_timestamp=TS,
        geom_geojson=SQUARE, registry_number="77:01-2.2",
    ))

    results = search_zones(db, registry_number_contains="59:00")
    assert len(results) == 1
    assert results[0].registry_number == "59:00-1.1"


def test_zones_to_excel_produces_real_readable_workbook(db):
    load_zone(db, ConstraintZone(
        zone_type="ООПТ", status=STATUS_OFFICIAL, source_name="НСПД (тест)", data_timestamp=TS,
        geom_geojson=SQUARE, registry_number="A1", document_basis="Приказ №1", regime="запрет строительства",
    ))
    zones = search_zones(db, zone_type="ООПТ")

    xlsx_bytes = zones_to_excel(zones)

    assert xlsx_bytes[:2] == b"PK"  # .xlsx - это zip-архив
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb.active
    header = [cell.value for cell in ws[1]]
    assert header == COLUMNS
    row2 = [cell.value for cell in ws[2]]
    assert row2[0] == "ООПТ"
    assert row2[2] == "A1"
    assert row2[3] == "запрет строительства"
    assert row2[4] == "Приказ №1"


def test_zones_to_excel_empty_list_still_has_header():
    xlsx_bytes = zones_to_excel([])
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb.active
    assert [cell.value for cell in ws[1]] == COLUMNS
    assert ws.max_row == 1
